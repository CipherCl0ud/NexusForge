# ── NVIDIA DLL pre-loader ────────────────────────────────────────────────────
# Must be the FIRST thing that executes — before rembg, onnxruntime, or any
# nvidia package is imported. This ensures all math and neural network libraries
# are locked into Windows memory before the AI models request them.
import os
import sys
import ctypes

def _preload_all_nvidia_dlls():
    if sys.platform != "win32": return
    
    # 1. Find the nvidia folder in your virtual environment
    nvidia_root = next((os.path.join(p, "nvidia") for p in sys.path if os.path.isdir(os.path.join(p, "nvidia"))), None)
    if not nvidia_root: return

    # 2. Collect EVERY .dll file inside the nvidia folder recursively
    all_dlls = []
    dll_folders = set()
    for root, _, files in os.walk(nvidia_root):
        for file in files:
            if file.lower().endswith(".dll"):
                all_dlls.append(os.path.join(root, file))
                dll_folders.add(root)

    # 3. Add to Windows PATH and DLL search directories
    for folder in dll_folders:
        os.environ["PATH"] = folder + os.pathsep + os.environ.get("PATH", "")
        try: os.add_dll_directory(folder)
        except Exception: pass

    # 4. Multi-pass loading to resolve dependency order automatically
    loaded = set()
    while True:
        progress = False
        for dll_path in all_dlls:
            if dll_path not in loaded:
                try:
                    ctypes.WinDLL(dll_path)
                    loaded.add(dll_path)
                    progress = True
                except Exception:
                    # Fails if a dependency isn't loaded yet; we will retry next loop
                    pass
        
        # If we didn't manage to load any new DLLs in this pass, break the loop
        if not progress:
            break
            
    print(f"[CUDA] Successfully locked {len(loaded)} NVIDIA DLL(s) into memory!")

_preload_all_nvidia_dlls()
# ─────────────────────────────────────────────────────────────────────────────

import io
import cv2
import gpxpy
import shutil
from datetime import datetime, timedelta
from fastapi import BackgroundTasks
from fastapi.responses import FileResponse
import numpy as np
import requests
import zipfile
import pymupdf as fitz
from pypdf import PdfReader, PdfWriter
from fastapi import FastAPI, File, UploadFile, Response, Form
from fastapi.middleware.cors import CORSMiddleware
from contextlib import asynccontextmanager
from rembg import remove
import onnxruntime as ort
import mammoth
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
import openpyxl
from docx import Document as DocxDocument
from typing import List
import tempfile
from PIL import Image
from moviepy.video.io.VideoFileClip import VideoFileClip

# ─── AI MODEL CONFIGURATION ──────────────────────────────────────────────────
MODEL_DIR  = "models"
MODEL_NAME = "RealESRGAN_x4.onnx"
MODEL_PATH = os.path.join(MODEL_DIR, MODEL_NAME)
MODEL_URL  = "https://huggingface.co/notaneimu/onnx-image-models/resolve/main/RealESRGAN_x4plus.onnx"

upscale_session = None

# ─── LIFESPAN (BOOT SEQUENCE) ────────────────────────────────────────────────
@asynccontextmanager
async def lifespan(app: FastAPI):
    global upscale_session
    print("Starting NexusForge Vision & Document Engine...")
    os.makedirs(MODEL_DIR, exist_ok=True)

    if not os.path.exists(MODEL_PATH) or os.path.getsize(MODEL_PATH) < 10_000:
        print("Downloading Real-ESRGAN AI Upscaling Model (~67MB)...")
        
        # Spoof a standard web browser to bypass Hugging Face script-blockers
        headers = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64)"}
        resp = requests.get(MODEL_URL, headers=headers, stream=True, timeout=60)
        
        if resp.status_code == 200:
            with open(MODEL_PATH, "wb") as f:
                for chunk in resp.iter_content(chunk_size=8192):
                    f.write(chunk)
            print("Model downloaded successfully!")
        else:
            raise RuntimeError(f"Model download failed — HTTP {resp.status_code}")

    print("Loading Real-ESRGAN into ONNX Runtime (GPU)...")
    upscale_session = ort.InferenceSession(
        MODEL_PATH,
        providers=['CUDAExecutionProvider', 'CPUExecutionProvider']
    )

    print("NexusForge Engine is LIVE on port 8000! 🚀")
    yield
    print("Shutting down NexusForge Engine...")

# ─── APP CONFIG & CORS ───────────────────────────────────────────────────────
app = FastAPI(lifespan=lifespan)

app.add_middleware(
    CORSMiddleware,
    allow_origins=["http://localhost:5173"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# ─── ENDPOINTS ───────────────────────────────────────────────────────────────

@app.post("/api/remove-bg")
async def remove_background(file: UploadFile = File(...)):
    try:
        # rembg automatically hooks into the pre-loaded ONNX GPU runtime
        result = remove(await file.read())
        return Response(content=result, media_type="image/png")
    except Exception as e:
        return {"error": str(e)}


@app.post("/api/upscale")
async def upscale_image(file: UploadFile = File(...)):
    try:
        # 1. Load image via OpenCV
        nparr = np.frombuffer(await file.read(), np.uint8)
        img = cv2.imdecode(nparr, cv2.IMREAD_COLOR)

        # 2. Pre-process for Real-ESRGAN (BGR -> RGB -> Float32 -> CHW format)
        img_rgb = cv2.cvtColor(img, cv2.COLOR_BGR2RGB)
        img_normalized = img_rgb.astype(np.float32) / 255.0
        img_transposed = np.transpose(img_normalized, (2, 0, 1))
        input_tensor = np.expand_dims(img_transposed, axis=0)  # Add batch dimension

        # 3. Run ONNX Inference on the GTX 1650
        input_name = upscale_session.get_inputs()[0].name
        output_tensor = upscale_session.run(None, {input_name: input_tensor})[0]

        # 4. Post-process (Remove Batch -> CHW -> HWC -> RGB -> BGR -> UInt8)
        output_image = np.squeeze(output_tensor, axis=0)
        output_image = np.clip(output_image, 0.0, 1.0)
        output_image = np.transpose(output_image, (1, 2, 0))
        output_image = (output_image * 255.0).round().astype(np.uint8)
        output_bgr = cv2.cvtColor(output_image, cv2.COLOR_RGB2BGR)

        # 5. Encode to PNG and Return
        _, buf = cv2.imencode(".png", output_bgr)
        return Response(content=buf.tobytes(), media_type="image/png")
        
    except Exception as e:
        print(f"Upscale Error: {e}")
        return {"error": str(e)}


@app.post("/api/compress-pdf")
async def compress_pdf(file: UploadFile = File(...)):
    try:
        raw = await file.read()

        # Pass 1: image quality reduction
        reader = PdfReader(io.BytesIO(raw))
        writer = PdfWriter()
        for page in reader.pages:
            writer.add_page(page)
        for page in writer.pages:
            for img in page.images:
                try:
                    img.replace(img.image, quality=40)
                except Exception:
                    pass
        buf1 = io.BytesIO()
        writer.write(buf1)

        # Pass 2: structural deflation
        doc = fitz.open(stream=buf1.getvalue(), filetype="pdf")
        out = doc.write(garbage=4, deflate=True, clean=True)
        return Response(content=out, media_type="application/pdf")
    except Exception as e:
        print(f"Compression error: {e}")
        return {"error": str(e)}


@app.post("/api/protect-pdf")
async def protect_pdf(
    file:     UploadFile = File(...),
    password: str        = Form(...),
    action:   str        = Form(...),
):
    try:
        doc = fitz.open(stream=await file.read(), filetype="pdf")
        if action == "encrypt":
            out = doc.write(
                encryption=fitz.PDF_ENCRYPT_AES_256,
                owner_pw=password,
                user_pw=password,
                permissions=fitz.PDF_PERM_PRINT | fitz.PDF_PERM_COPY,
            )
        elif action == "decrypt":
            if doc.needs_pass and not doc.authenticate(password):
                return {"error": "Incorrect password"}
            out = doc.write(encryption=fitz.PDF_ENCRYPT_NONE)
        else:
            return {"error": f"Unknown action: {action}"}
        return Response(content=out, media_type="application/pdf")
    except Exception as e:
        return {"error": str(e)}


@app.post("/api/pdf-to-image")
async def pdf_to_image(file: UploadFile = File(...)):
    try:
        doc     = fitz.open(stream=await file.read(), filetype="pdf")
        zip_buf = io.BytesIO()
        with zipfile.ZipFile(zip_buf, "a", zipfile.ZIP_DEFLATED, False) as zf:
            for i in range(len(doc)):
                pix = doc.load_page(i).get_pixmap(matrix=fitz.Matrix(2, 2))
                zf.writestr(f"page_{i + 1}.png", pix.tobytes("png"))
        return Response(content=zip_buf.getvalue(), media_type="application/zip")
    except Exception as e:
        return {"error": str(e)}
    
EMU_PER_POINT = 12700  # EMU (pptx units) -> PDF points
 
 
def emu_to_pt(emu):
    return (emu or 0) / EMU_PER_POINT
 
 
def html_to_pdf_bytes(html: str) -> bytes:
    """Renders HTML/CSS to a PDF using PyMuPDF's built-in Story engine (requires PyMuPDF >= 1.21)."""
    story = fitz.Story(html=html)
    buf = io.BytesIO()
    writer = fitz.DocumentWriter(buf)
    mediabox = fitz.paper_rect("a4")
    where = mediabox + (36, 36, -36, -36)  # margins
    more = True
    while more:
        dev = writer.begin_page(mediabox)
        more, _ = story.place(where)
        story.draw(dev)
        writer.end_page()
    writer.close()
    return buf.getvalue()
 
 
def package_results(results: list, single_media_type: str):
    """results: list of (filename, bytes). Returns the file directly if there's one result, else a zip."""
    if len(results) == 1:
        return Response(content=results[0][1], media_type=single_media_type)
    zip_buf = io.BytesIO()
    with zipfile.ZipFile(zip_buf, "a", zipfile.ZIP_DEFLATED, False) as zf:
        for name, data in results:
            zf.writestr(name, data)
    return Response(content=zip_buf.getvalue(), media_type="application/zip")
 
 
# ─── WORD → PDF ───────────────────────────────────────────────────────────────
@app.post("/api/word-to-pdf")
async def word_to_pdf(files: List[UploadFile] = File(...)):
    try:
        results = []
        for f in files:
            raw = await f.read()
            html_result = mammoth.convert_to_html(io.BytesIO(raw))
            styled_html = f"""
            <html><head><style>
                body {{ font-family: Helvetica, Arial, sans-serif; font-size: 11pt; line-height: 1.5; }}
                h1, h2, h3 {{ font-family: Helvetica, Arial, sans-serif; }}
                img {{ max-width: 100%; }}
            </style></head><body>{html_result.value}</body></html>
            """
            pdf_bytes = html_to_pdf_bytes(styled_html)
            results.append((os.path.splitext(f.filename)[0] + ".pdf", pdf_bytes))
 
        return package_results(results, "application/pdf")
    except Exception as e:
        print(f"Word to PDF error: {e}")
        return {"error": str(e)}
 
 
# ─── PPT → PDF ────────────────────────────────────────────────────────────────
@app.post("/api/ppt-to-pdf")
async def ppt_to_pdf(files: List[UploadFile] = File(...)):
    try:
        results = []
        for f in files:
            raw = await f.read()
            prs = Presentation(io.BytesIO(raw))
            page_w = emu_to_pt(prs.slide_width)
            page_h = emu_to_pt(prs.slide_height)
 
            doc = fitz.open()
            for slide in prs.slides:
                page = doc.new_page(width=page_w, height=page_h)
                for shape in slide.shapes:
                    try:
                        left, top = emu_to_pt(shape.left), emu_to_pt(shape.top)
                        width, height = emu_to_pt(shape.width), emu_to_pt(shape.height)
                        rect = fitz.Rect(left, top, left + width, top + height)
 
                        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                            page.insert_image(rect, stream=shape.image.blob)
                        elif shape.has_text_frame and shape.text_frame.text.strip():
                            fontsize = 14
                            try:
                                run = shape.text_frame.paragraphs[0].runs[0]
                                if run.font.size:
                                    fontsize = run.font.size.pt
                            except Exception:
                                pass
                            page.insert_textbox(rect, shape.text_frame.text, fontsize=fontsize, fontname="helv")
                    except Exception as shape_err:
                        print(f"Shape skipped: {shape_err}")
                        continue
 
            out_buf = io.BytesIO()
            doc.save(out_buf)
            results.append((os.path.splitext(f.filename)[0] + ".pdf", out_buf.getvalue()))
 
        return package_results(results, "application/pdf")
    except Exception as e:
        print(f"PPT to PDF error: {e}")
        return {"error": str(e)}
 
 
# ─── EXCEL → PDF ──────────────────────────────────────────────────────────────
@app.post("/api/excel-to-pdf")
async def excel_to_pdf(files: List[UploadFile] = File(...)):
    try:
        results = []
        for f in files:
            raw = await f.read()
            wb = openpyxl.load_workbook(io.BytesIO(raw), data_only=True)
 
            html_parts = [
                "<html><head><style>",
                "body{font-family:Helvetica,Arial,sans-serif;font-size:9pt;}",
                "table{border-collapse:collapse;width:100%;margin-bottom:20px;}",
                "td,th{border:1px solid #999;padding:4px 6px;text-align:left;}",
                "h2{font-family:Helvetica,Arial,sans-serif;}",
                "</style></head><body>",
            ]
            for sheet in wb.worksheets:
                html_parts.append(f"<h2>{sheet.title}</h2><table>")
                for row in sheet.iter_rows():
                    html_parts.append("<tr>")
                    for cell in row:
                        val = "" if cell.value is None else str(cell.value)
                        html_parts.append(f"<td>{val}</td>")
                    html_parts.append("</tr>")
                html_parts.append("</table>")
            html_parts.append("</body></html>")
 
            pdf_bytes = html_to_pdf_bytes("".join(html_parts))
            results.append((os.path.splitext(f.filename)[0] + ".pdf", pdf_bytes))
 
        return package_results(results, "application/pdf")
    except Exception as e:
        print(f"Excel to PDF error: {e}")
        return {"error": str(e)}
 
 
# ─── PDF → WORD ───────────────────────────────────────────────────────────────
@app.post("/api/pdf-to-word")
async def pdf_to_word(files: List[UploadFile] = File(...)):
    """Text-only extraction — layout, images, and tables are not preserved."""
    try:
        results = []
        for f in files:
            raw = await f.read()
            pdf = fitz.open(stream=raw, filetype="pdf")
            document = DocxDocument()
            for i, page in enumerate(pdf):
                if i > 0:
                    document.add_page_break()
                text = page.get_text("text")
                for para in text.split("\n"):
                    if para.strip():
                        document.add_paragraph(para)
 
            out_buf = io.BytesIO()
            document.save(out_buf)
            results.append((os.path.splitext(f.filename)[0] + ".docx", out_buf.getvalue()))
 
        return package_results(
            results, "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )
    except Exception as e:
        print(f"PDF to Word error: {e}")
        return {"error": str(e)}
    # ─── VIDEO TO GIF & AUDIO EXTRACTOR ──────────────────────────────────────────

@app.post("/api/video-to-gif")
async def video_to_gif(
    file: UploadFile = File(...),
    fps: int = Form(10),
    width: int = Form(480),
    speed: float = Form(1.0)
):
    """Converts uploaded video into an optimized animated GIF."""
    suffix = os.path.splitext(file.filename)[1] or ".mp4"
    with tempfile.NamedTemporaryFile(delete=False, suffix=suffix) as temp_in:
        temp_in.write(await file.read())
        temp_in_path = temp_in.name

    try:
        cap = cv2.VideoCapture(temp_in_path)
        video_fps = cap.get(cv2.CAP_PROP_FPS) or 30.0
        frame_step = max(1, int(round((video_fps / fps) / speed)))

        frames = []
        frame_idx = 0

        while cap.isOpened():
            ret, frame = cap.read()
            if not ret:
                break

            if frame_idx % frame_step == 0:
                h, w = frame.shape[:2]
                target_w = min(width, w)
                target_h = int(h * (target_w / w))
                
                resized = cv2.resize(frame, (target_w, target_h), interpolation=cv2.INTER_AREA)
                rgb_frame = cv2.cvtColor(resized, cv2.COLOR_BGR2RGB)
                
                # Convert to PIL and apply color quantization for compact GIF size
                pil_img = Image.fromarray(rgb_frame).convert("P", palette=Image.ADAPTIVE, colors=128)
                frames.append(pil_img)

            frame_idx += 1

        cap.release()

        if not frames:
            raise ValueError("No frames could be extracted from video.")

        gif_buffer = io.BytesIO()
        duration_ms = int(1000 / fps)
        
        frames[0].save(
            gif_buffer,
            format="GIF",
            save_all=True,
            append_images=frames[1:],
            duration=duration_ms,
            loop=0,
            optimize=True
        )

        out_name = os.path.splitext(file.filename)[0] + ".gif"
        return Response(
            content=gif_buffer.getvalue(),
            media_type="image/gif",
            headers={"Content-Disposition": f'attachment; filename="{out_name}"'}
        )
    finally:
        if os.path.exists(temp_in_path):
            os.remove(temp_in_path)


@app.post("/api/audio-extract")
async def audio_extract(
    file: UploadFile = File(...),
    format: str = Form("mp3")
):
    """Extracts audio stream from video and outputs MP3 or WAV."""
    suffix = os.path.splitext(file.filename)[1] or ".mp4"
    with tempfile.NamedTemporaryFile(delete=False, suffix=suffix) as temp_in:
        temp_in.write(await file.read())
        temp_in_path = temp_in.name

    out_ext = ".mp3" if format.lower() == "mp3" else ".wav"
    temp_out_path = temp_in_path + out_ext

    try:
        clip = VideoFileClip(temp_in_path)
        if clip.audio is None:
            raise ValueError("Selected video has no audio track.")

        codec = "libmp3lame" if format.lower() == "mp3" else "pcm_s16le"
        clip.audio.write_audiofile(temp_out_path, codec=codec, logger=None)
        clip.close()

        with open(temp_out_path, "rb") as f:
            audio_bytes = f.read()

        media_type = "audio/mpeg" if format.lower() == "mp3" else "audio/wav"
        out_name = os.path.splitext(file.filename)[0] + out_ext

        return Response(
            content=audio_bytes,
            media_type=media_type,
            headers={"Content-Disposition": f'attachment; filename="{out_name}"'}
        )
    finally:
        if os.path.exists(temp_in_path):
            os.remove(temp_in_path)
        if os.path.exists(temp_out_path):
            os.remove(temp_out_path)
# ─── VIDEO & GPX PROCESSING CLASSES ──────────────────────────────────────────

class Coordinate_Systems:
    @staticmethod
    def generate_json_gpx(gpx):
        gpx_json = []
        idx = 0
        for track in gpx.tracks:
            for segment in track.segments:
                for point in segment.points:
                    time_iso = point.time.isoformat() if point.time else None
                    gpx_json.append({idx: {"long": point.longitude, "lat": point.latitude, "time": time_iso}})
                    idx += 1
        return gpx_json
    
    @staticmethod
    def average_json(gpx_json):
        parsed_data = []
        for entry in gpx_json:
            key, value = list(entry.items())[0]
            value['index'] = int(key)  
            value['time'] = datetime.fromisoformat(value['time'])
            parsed_data.append(value)

        result = []
        new_index = 0  
        for i in range(len(parsed_data) - 1):
            current = parsed_data[i]
            next_entry = parsed_data[i + 1]
            
            result.append({str(new_index): {"long": current["long"], "lat": current["lat"], "time": current["time"].isoformat()}})
            new_index += 1
            
            current_time = current["time"]
            next_time = next_entry["time"]
            while (current_time + timedelta(seconds=1)) < next_time:
                current_time += timedelta(seconds=1)
                result.append({
                    str(new_index): {
                        "long": (current["long"] + next_entry["long"]) / 2,
                        "lat": (current["lat"] + next_entry["lat"]) / 2,
                        "time": current_time.isoformat()
                    }
                })
                new_index += 1

        last_entry = parsed_data[-1]
        result.append({str(new_index): {"long": last_entry["long"], "lat": last_entry["lat"], "time": last_entry["time"].isoformat()}})
        return result
    
    @staticmethod
    def get_gpx_coords(gpx_path: str, average: bool = False):
        with open(gpx_path, 'r') as gpx_file:
            gpx = gpxpy.parse(gpx_file)
        gpx_json = Coordinate_Systems.generate_json_gpx(gpx)
        if average:
            gpx_json = Coordinate_Systems.average_json(gpx_json) 
        return gpx_json
    
    @staticmethod
    def location_from_gpx(gpx_json, index):
        for idxs in gpx_json:
            for key, val in idxs.items():
                if int(key) == int(index):
                    return [val["long"], val["lat"]]
        return []
    
    @staticmethod
    def point_json_snippet(coord_list, file_path, time_sec):
        return {
            "type": "Feature",
            "properties": {
                "file_name": os.path.basename(file_path),
                "time_in_sec": time_sec,
            },
            "geometry": {
                "coordinates": coord_list,
                "type": "Point"
            }
        }

class Video2ImageProcessor:
    def __init__(self, video_path, gpx_path, start_time_sec, end_time_sec, interval, average_gpx, output_dir):
        self.video_path = video_path
        self.gpx_path = gpx_path
        self.start_time_sec = start_time_sec
        self.end_time_sec = end_time_sec
        self.interval = interval
        self.average_gpx = average_gpx
        self.output_dir = output_dir

    def process(self):
        gpx_json = Coordinate_Systems.get_gpx_coords(self.gpx_path, self.average_gpx)
        cap = cv2.VideoCapture(self.video_path)
        output_fps = int(cap.get(cv2.CAP_PROP_FPS))
        frame_count = int(cap.get(cv2.CAP_PROP_FRAME_COUNT))
        duration = frame_count / output_fps if output_fps > 0 else 0

        if self.start_time_sec > duration or (self.end_time_sec and self.end_time_sec > duration):
            cap.release()
            raise ValueError("Time bounds exceed video duration.")
        
        actual_end = self.end_time_sec if self.end_time_sec else int(duration)
        frame_interval = int(output_fps * self.interval)
        frame_idx = 0
        
        json_features = []
        video_name = os.path.splitext(os.path.basename(self.video_path))[0]
        
        while True:
            ret, frame = cap.read()
            if not ret: break
            
            if (self.start_time_sec * output_fps) <= frame_idx <= (actual_end * output_fps):
                secs = frame_idx % frame_interval
                if secs == 0:
                    loc = Coordinate_Systems.location_from_gpx(gpx_json, secs)
                    if len(loc) == 2:
                        frame_filename = os.path.join(self.output_dir, f'{video_name}_{frame_idx}.jpg')
                        cv2.imwrite(frame_filename, frame)
                        json_features.append(Coordinate_Systems.point_json_snippet(loc, frame_filename, secs))
            frame_idx += 1
        cap.release()

        geojson_data = {"type": "FeatureCollection", "features": json_features}
        geojson_path = os.path.join(self.output_dir, f"{video_name}.geojson")
        import json
        with open(geojson_path, "w") as f:
            json.dump(geojson_data, f, indent=2)

        zip_path = f"{self.output_dir}.zip"
        with zipfile.ZipFile(zip_path, 'w') as zipf:
            for root, _, files in os.walk(self.output_dir):
                for file in files:
                    zipf.write(os.path.join(root, file), file)
                    
        return zip_path

# ─── VIDEO ENDPOINTS ─────────────────────────────────────────────────────────

def cleanup_temp_files(paths: list):
    """Background task to delete temp folders and zips after sending to client"""
    for path in paths:
        try:
            if os.path.isdir(path): shutil.rmtree(path)
            elif os.path.isfile(path): os.remove(path)
        except Exception as e:
            print(f"Cleanup failed for {path}: {e}")

def time_to_seconds(time_str: str) -> int:
    time_str = time_str.replace(".", ":")
    parts = list(map(int, time_str.split(':')))
    if len(parts) == 3: return parts[0] * 3600 + parts[1] * 60 + parts[2]
    elif len(parts) == 2: return parts[0] * 60 + parts[1]
    return parts[0]
# ─────────────────────────────────────────────────────────────────────────────
# VIDEO TO IMAGES (no GPX) — paste into main.py, below the existing
# Video2ImageProcessor class / video endpoints. Reuses time_to_seconds() and
# cleanup_temp_files(), which already exist in main.py. No new pip installs.
# ─────────────────────────────────────────────────────────────────────────────

class VideoFrameExtractor:
    """Extracts frames from a video at a fixed time interval — no GPS tagging."""

    def __init__(self, video_path, start_time_sec, end_time_sec, interval, output_dir):
        self.video_path = video_path
        self.start_time_sec = start_time_sec
        self.end_time_sec = end_time_sec
        self.interval = interval
        self.output_dir = output_dir

    def process(self):
        cap = cv2.VideoCapture(self.video_path)
        fps = cap.get(cv2.CAP_PROP_FPS) or 0
        frame_count = int(cap.get(cv2.CAP_PROP_FRAME_COUNT))
        duration = frame_count / fps if fps > 0 else 0

        if self.start_time_sec > duration or (self.end_time_sec and self.end_time_sec > duration):
            cap.release()
            raise ValueError("Time bounds exceed video duration.")

        actual_end = self.end_time_sec if self.end_time_sec else duration
        video_name = os.path.splitext(os.path.basename(self.video_path))[0]

        frame_idx = 0
        saved = 0
        next_capture_time = self.start_time_sec

        while True:
            ret, frame = cap.read()
            if not ret:
                break

            current_time = (frame_idx / fps) if fps > 0 else 0

            if self.start_time_sec <= current_time <= actual_end and current_time >= next_capture_time:
                frame_filename = os.path.join(
                    self.output_dir, f"{video_name}_{saved:04d}_{current_time:.2f}s.jpg"
                )
                cv2.imwrite(frame_filename, frame)
                saved += 1
                next_capture_time += self.interval

            frame_idx += 1

        cap.release()

        zip_path = f"{self.output_dir}.zip"
        with zipfile.ZipFile(zip_path, "w") as zipf:
            for root, _, files in os.walk(self.output_dir):
                for file in files:
                    zipf.write(os.path.join(root, file), file)

        return zip_path, saved


@app.post("/api/video-to-frames")
async def extract_video_frames(
    background_tasks: BackgroundTasks,
    video: UploadFile = File(...),
    start_time: str = Form("0:0:0"),
    end_time: str = Form(""),
    interval: float = Form(1.0),
):
    if not video.filename.lower().endswith(".mp4"):
        return {"error": "Only MP4 files are accepted."}

    session_id = datetime.now().strftime("%Y%m%d%H%M%S")
    temp_dir = os.path.join("temp_data", session_id)
    frames_dir = os.path.join(temp_dir, "frames")
    os.makedirs(frames_dir, exist_ok=True)

    vid_path = os.path.join(temp_dir, video.filename)
    with open(vid_path, "wb") as f:
        f.write(await video.read())

    start_sec = time_to_seconds(start_time)
    end_sec = time_to_seconds(end_time) if end_time else None

    try:
        extractor = VideoFrameExtractor(
            video_path=vid_path,
            start_time_sec=start_sec,
            end_time_sec=end_sec,
            interval=interval,
            output_dir=frames_dir,
        )
        zip_path, saved_count = extractor.process()

        if saved_count == 0:
            cleanup_temp_files([temp_dir])
            return {"error": "No frames were extracted — check your time range and interval."}

        background_tasks.add_task(cleanup_temp_files, [temp_dir, zip_path])

        return FileResponse(
            zip_path, media_type="application/zip", filename=f"{session_id}_frames.zip"
        )

    except Exception as e:
        print(f"Frame extraction error: {e}")
        cleanup_temp_files([temp_dir])
        return {"error": str(e)}

@app.post("/api/video-to-gpx-frames")
async def extract_video_gpx_frames(
    background_tasks: BackgroundTasks,
    video: UploadFile = File(...),
    gpx: UploadFile = File(...),
    start_time: str = Form("0:0:0"),
    end_time: str = Form(""),
    interval: float = Form(1.0),
    gpx_correction: bool = Form(False)
):
    if not video.filename.lower().endswith(".mp4"):
        return {"error": "Only MP4 files are accepted."}
    if not gpx.filename.lower().endswith(".gpx"):
        return {"error": "Input GPX file is invalid."}

    # Setup temp directories
    session_id = datetime.now().strftime("%Y%m%d%H%M%S")
    temp_dir = os.path.join("temp_data", session_id)
    frames_dir = os.path.join(temp_dir, "frames")
    os.makedirs(frames_dir, exist_ok=True)
    
    vid_path = os.path.join(temp_dir, video.filename)
    gpx_path = os.path.join(temp_dir, gpx.filename)
    
    # Save uploads to disk
    with open(vid_path, "wb") as f: f.write(await video.read())
    with open(gpx_path, "wb") as f: f.write(await gpx.read())
    
    start_sec = time_to_seconds(start_time)
    end_sec = time_to_seconds(end_time) if end_time else None

    try:
        processor = Video2ImageProcessor(
            video_path=vid_path,
            gpx_path=gpx_path,
            start_time_sec=start_sec,
            end_time_sec=end_sec,
            interval=interval,
            average_gpx=gpx_correction,
            output_dir=frames_dir
        )
        zip_path = processor.process()
        
        # Schedule cleanup to run AFTER the ZIP file is transmitted to the user
        background_tasks.add_task(cleanup_temp_files, [temp_dir, zip_path])
        
        return FileResponse(zip_path, media_type="application/zip", filename=f"{session_id}_frames.zip")
        
    except Exception as e:
        cleanup_temp_files([temp_dir])
        return {"error": str(e)}

if __name__ == "__main__":
    import uvicorn
    uvicorn.run(app, host="0.0.0.0", port=8000)